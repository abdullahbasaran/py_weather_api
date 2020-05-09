import requests
import pprint
import pandas as pd
from unidecode import unidecode
import schedule
import time
from pandas import ExcelWriter
from pandas import ExcelFile
import numpy as np
from datetime import date, timedelta
import datetime
from openpyxl import load_workbook
import json

# excel den dataframe e donusturme
excel_data_df = pd.read_excel('ilce_merkez_coord.xlsx')

for col in ['ILCEADI']:
    excel_data_df[col] = excel_data_df[col].apply(
        unidecode)  # turkce karakterler decode edildi

print(excel_data_df)  # turkce karakterlerden bagimsiz terminale respose yazilmasi

i = 1
# edate = date(2008, 5, 5)   # end date

dt = date(2020, 3, 1)
end = date.today()
step = datetime.timedelta(days=1)

'''
book = load_workbook('New_Excel.xlsx')
writer = pd.ExcelWriter('New_Excel.xlsx', engine='openpyxl')
writer.book = book
'''


def WriteToExcel(kod_ilce,sdate,temp,sunrise,sunset,moonrise,moonset, hour, wind_speed, temperature):

    hourly_label= ""
    hourInt = int(hour)

    if(hourInt == 0):
        hourly_label = "(00:00-03:00)"
    elif(hourInt == 300):
        hourly_label = "(03:00-06:00) "
        
    elif(hourInt == 600):
        hourly_label = "(06:00-09:00) "
        
    elif(hourInt == 900):
        hourly_label = "(09:00-12:00) "
        
    elif(hourInt == 1200):
        hourly_label = "(12:00-15:00) "
        
    elif(hourInt == 1500):
        hourly_label = "(15:00-18:00) "
        
    elif(hourInt == 1800):
        hourly_label = "(18:00-21:00) "
        
    else :
        hourly_label = "(21:00-00:00) "
    

    df = pd.DataFrame({'ILCEKOD': [kod_ilce], 'TIME': [sdate], 'TEMPRATURE': [temp], 'SUNRISE': [sunrise], 'SUNSET': [
                      sunset], 'MOONRISE': [moonrise], 'MOONSET': [moonset], 'HOURLY': [hourly_label], 'WINDSPEED': [wind_speed], 'TEMPERATURE': [temperature]})  # excel tablosunda her bir sutun icin degerlerin  girilmesi
    writer = pd.ExcelWriter('Excel2.xlsx',mode='')
    df.to_excel(writer, 'Sheet1', index=False)
    # writer.save()
    # try to open an existing workbook
    writer.book = load_workbook('Excel2.xlsx')

    # copy existing sheets
    writer.sheets = dict((ws.title, ws) for ws in writer.book.worksheets)
    # read existing file

    reader = pd.read_excel(r'Excel2.xlsx')
    # write out the new sheet
    df.to_excel(writer, index=False, header=False, startrow=len(reader)+1)
    writer.close()


while dt < end:

    sdate = str(dt.strftime('%Y-%m-%d'))
    # http://api.weatherstack.com/historical?access_key=51d7bfab86b02431a6a6cc5d489ea60b&query=Aladag&historical_date=2015-01-21&hourly=1
    ilce = excel_data_df.loc[0][5]
    kod_ilce = excel_data_df.loc[0][3]
    url = 'https://api.weatherstack.com/historical?access_key=51d7bfab86b02431a6a6cc5d489ea60b&hourly=1&historical_date='+sdate+'&query='+ilce
    res = requests.get(url)
    data = res.json()
    # time = data['location']['localtime'] #istenilen degerin cekilmesi
    temp = data['current']['temperature']  # istenilen degerin cekilmesi
    sunrise = data['historical'][sdate]['astro']['sunrise']
    sunset = data['historical'][sdate]['astro']['sunset']
    moonrise = data['historical'][sdate]['astro']['moonrise']
    moonset = data['historical'][sdate]['astro']['moonset']
    moon_illumination = data['historical'][sdate]['astro']['moon_illumination']

    for hourly_object in data['historical'][sdate]['hourly']:
        WriteToExcel(kod_ilce,sdate,temp,sunrise,sunset,moonrise,moonset,hourly_object['time'], hourly_object['wind_speed'], hourly_object['temperature'])
    
    dt += step
    i = i+1
''' 


writer = ExcelWriter('Excel2.xlsx',mode='a')
df.to_excel(writer,'Sheet1',index=False)
#max = max_row
#max=+1
writer.save()

writer.sheets=dict((ws.title,ws) for ws in book.worksheets)
df.to_excel(writer,'Sheet1',startrow=i, index=False)

max = ws.max_row
for row, entry in enumerate(data1,start=1):
st.cell(row=row+max, column=1, value=entry)

writer.save()
'''


# schedule.every(10).seconds.do(WriteToExcel) #excele schedule ile yazma
# while True:
#    schedule.run_pending()
# time.sleep(10)


# print(data)
